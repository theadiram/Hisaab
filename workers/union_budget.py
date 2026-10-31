"""Private, PDF-checked Union budget context. Never work/payment events."""

import hashlib, json, re, urllib.request
from io import BytesIO
from decimal import Decimal
from pathlib import Path
import openpyxl
from pypdf import PdfReader
from ingest import DATA, ACQUISITION, ROOT, atomic_json, now

BASE = "https://www.indiabudget.gov.in/"
SOURCE = "union-budget-2026"
VERSION = "union-demands-pdf-checked-v1"
FILES = {
    "summary.xlsx": "doc/eb/sumsbe.xlsx",
    "summary.pdf": "doc/eb/sumsbe.pdf",
    "detail.xlsx": "doc/eb/allsbe.xlsx",
    "detail.pdf": "doc/eb/allsbe.pdf",
}
PERIODS = [
    ("actual", "2024-2025"),
    ("budget_estimate", "2025-2026"),
    ("revised_estimate", "2025-2026"),
    ("budget_estimate", "2026-2027"),
]


def amount(value):
    if value == "...":
        return None
    if (
        value is None
        or isinstance(value, bool)
        or (isinstance(value, str) and value.startswith("="))
    ):
        raise ValueError("Missing/formula amount")
    d = Decimal(str(value).replace(",", ""))
    if not d.is_finite() or d != d.quantize(Decimal(".01")):
        raise ValueError("Invalid amount precision")
    return format(d, ".2f")


def parse_summary(rows):
    if not any("In ₹ Crores" in str(x) for row in rows[:5] for x in row):
        raise ValueError("Summary currency unit changed")
    result = {}
    for row_number, row in enumerate(rows, 1):
        if len(row) < 7:
            continue
        m = re.fullmatch(r"\s*(\d+)\.\s*", str(row[1]))
        if not m:
            continue
        number = int(m[1])
        if number in result or not row[2]:
            raise ValueError("Duplicate/missing demand identity")
        values = [amount(x) for x in row[3:6]]
        if all(x is not None for x in values) and Decimal(values[0]) + Decimal(
            values[1]
        ) != Decimal(values[2]):
            raise ValueError("Summary does not reconcile")
        result[number] = {"department": row[2], "values": values, "row": row_number}
    if set(result) != set(range(1, 103)):
        raise ValueError("Expected complete 102-demand edition")
    return result


def pdf_net_rows(reader):
    result = {}
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text()
        identity = re.search(r"\bDEMAND\s+NO\.\s*(\d+)\b", text) or re.search(
            r"No\.\s*(\d+)\s*\(APPROPRIATION\)", text
        )
        if not identity:
            continue
        number = int(identity[1])
        if number in result:
            continue
        lines = [line.strip() for line in text.splitlines() if re.match(r"^\s*Net\s", line)]
        if not lines:
            raise ValueError("PDF net row absent")
        tokens = lines[0].split()[1:]
        if len(tokens) != 12:
            raise ValueError("PDF net row layout changed")
        result[number] = {
            "values": [amount(v) for v in tokens],
            "page": page_number,
            "rawText": lines[0],
        }
    if set(result) != set(range(1, 103)):
        raise ValueError("Incomplete PDF demand identities")
    return result


def parse_demand(rows, number, summary, pdf):
    identity = "\n".join(str(v) for row in rows[:8] for v in row if v is not None)
    if (
        not (
            re.search(r"Demand No\.\s*" + str(number) + r"\b", identity)
            or re.search(r"No\.\s*" + str(number) + r"\s*\(APPROPRIATION\)", identity)
        )
        or "(In ₹ Crores)" not in identity
    ):
        raise ValueError("Demand identity/unit changed")
    period_row = next(
        (i for i, row in enumerate(rows[:15]) if any(str(v).startswith("Actuals ") for v in row)),
        None,
    )
    if period_row is None:
        raise ValueError("Period header missing")
    starts = [
        (i, str(v))
        for i, v in enumerate(rows[period_row])
        if re.fullmatch(r"(Actuals|Budget Estimates|Revised Estimates) \d{4}-\d{4}", str(v))
    ]
    expected = [
        "Actuals 2024-2025",
        "Budget Estimates 2025-2026",
        "Revised Estimates 2025-2026",
        "Budget Estimates 2026-2027",
    ]
    if [v for _, v in starts] != expected:
        raise ValueError("Unexpected period order")
    stop = next(
        (
            i
            for i, row in enumerate(rows)
            if any(str(v).startswith("A. The Budget allocations") for v in row)
        ),
        20,
    )
    candidates = [(i, row) for i, row in enumerate(rows[:stop]) if "Net" in row]
    if len(candidates) != 1:
        raise ValueError("Ambiguous demand net row")
    net_index, net = candidates[0]
    header = rows[period_row + 1]
    records = []
    flat = []
    for p, (start, _) in enumerate(starts):
        stop = starts[p + 1][0] if p + 1 < len(starts) else len(header)
        cols = [i for i in range(start, stop) if header[i] in ("Revenue", "Capital", "Total")]
        if [header[i] for i in cols] != ["Revenue", "Capital", "Total"]:
            raise ValueError("Ambiguous classification columns")
        values = [amount(net[i]) for i in cols]
        flat += values
        if all(v is not None for v in values) and Decimal(values[0]) + Decimal(
            values[1]
        ) != Decimal(values[2]):
            raise ValueError("Demand net row does not reconcile")
        measure, year = PERIODS[p]
        records.append(
            {
                "identity": f"{SOURCE}:demand:{number}:{measure}:{year}",
                "demandNumber": number,
                "department": summary["department"],
                "measure": measure,
                "financialYear": year,
                "revenueCrore": values[0],
                "capitalCrore": values[1],
                "totalCrore": values[2],
                "unit": "crore",
                "scope": "union_demand_net",
                "workId": None,
                "location": None,
                "workPayment": False,
                "locator": f"sbe{number}!row {net_index + 1}; authoritative PDF page {pdf['page']}",
                "pdfPage": pdf["page"],
            }
        )
    if flat != pdf["values"]:
        raise ValueError("Spreadsheet differs from authoritative PDF net row")
    differences = [(a, b) for a, b in zip(flat[-3:], summary["values"]) if a != b]
    if any(a is not None or b != "0.00" for a, b in differences):
        raise ValueError("Summary differs from detail net row")
    if differences:
        for record in records:
            record["notationNote"] = (
                "The authoritative PDF uses ellipsis where the summary workbook shows zero; ellipsis remains unspecified, not converted to zero."
            )
    return records


def run():
    documents = {}
    paths = {}
    retrieved = now()
    for name, relative in FILES.items():
        url = BASE + relative
        with urllib.request.urlopen(url, timeout=30) as response:
            if response.geturl() != url:
                raise ValueError("Unexpected source redirect")
            raw = response.read(50_000_001)
        if len(raw) > 50_000_000 or not raw.startswith(
            b"%PDF-" if name.endswith(".pdf") else b"PK"
        ):
            raise ValueError("Invalid original")
        digest = hashlib.sha256(raw).hexdigest()
        p = DATA / "originals" / SOURCE / digest
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(raw)
        paths[name] = p
        documents[name] = digest
        atomic_json(
            ACQUISITION / "documents" / f"{digest}.json",
            {
                "sourceId": SOURCE,
                "sourceName": "Union Budget 2026–27 official expenditure statements",
                "sourceFamily": "national_financial_context",
                "sha256": digest,
                "url": url,
                "contextUrl": BASE,
                "contentType": "application/pdf"
                if name.endswith(".pdf")
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "original": str(p.relative_to(ROOT)),
                "retrievedAt": retrieved,
                "parserVersion": VERSION,
                "permissionStatus": "permission_required",
                "permission": "Private assessment only. Ministry of Finance permission required by website copyright policy.",
                "scope": "Budget context, not project payment receipts.",
            },
        )
    summary = parse_summary(
        list(
            openpyxl.load_workbook(
                BytesIO(paths["summary.xlsx"].read_bytes()), read_only=True, data_only=False
            ).active.iter_rows(values_only=True)
        )
    )
    pdf = pdf_net_rows(PdfReader(paths["detail.pdf"]))
    workbook = openpyxl.load_workbook(
        BytesIO(paths["detail.xlsx"].read_bytes()), read_only=True, data_only=False
    )
    records = []
    quarantine = []
    for number in range(1, 103):
        try:
            records.extend(
                parse_demand(
                    list(workbook[f"sbe{number}"].iter_rows(values_only=True)),
                    number,
                    summary[number],
                    pdf[number],
                )
            )
        except (ValueError, KeyError) as e:
            quarantine.append({"demandNumber": number, "reason": str(e)})
    atomic_json(
        ACQUISITION / "union-budget-context.json",
        {
            "sourceId": SOURCE,
            "documents": documents,
            "records": records,
            "quarantine": quarantine,
            "parserVersion": VERSION,
        },
    )
    atomic_json(
        DATA / "union-budget-acquisition.json",
        {
            "sourceId": SOURCE,
            "retrievedAt": retrieved,
            "edition": "2026-2027",
            "demandStatements": 102,
            "pdfCheckedDemands": len(records) // 4,
            "checkedDemandNumbers": sorted({r["demandNumber"] for r in records}),
            "financialContextRows": len(records),
            "quarantinedDemands": len(quarantine),
            "periods": [{"measure": m, "financialYear": y} for m, y in PERIODS],
            "originalDocuments": 4,
            "hostedOriginalsVerified": False,
            "hostedContextRows": 0,
            "publication": "private_permission_and_review_pending",
            "scope": "Union demand totals; includes expenditure outside development. Not individual works or contractor payments.",
        },
    )
    print(
        json.dumps(
            {
                "pdfCheckedDemands": len(records) // 4,
                "contextRows": len(records),
                "quarantinedDemands": len(quarantine),
            }
        )
    )


if __name__ == "__main__":
    run()
